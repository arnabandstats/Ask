"""Shared fixtures: an isolated data dir, sample repos/docs/data, and a scripted fake LLM."""
from __future__ import annotations

import copy
import json
import os
import sys
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from ask import config  # noqa: E402
from ask.sources import loaders  # noqa: E402
from ask.sources.registry import SourceRegistry  # noqa: E402


# ── isolation ──────────────────────────────────────────────────────────────

@pytest.fixture(autouse=True)
def isolated_data_dir(tmp_path, monkeypatch):
    """Every test gets its own ask_data (chat DB, outputs, text cache)."""
    data = tmp_path / "ask_data"
    monkeypatch.setattr(config, "DATA_DIR", data)
    monkeypatch.setattr(config, "DB_PATH", data / "chats.db")
    monkeypatch.setattr(config, "OUTPUT_DIR", data / "outputs")
    loaders._TEXT_CACHE.clear()
    from ask.memory import store
    store.use_folder(None)                 # start every test at the default chat history folder
    store.set_resolver(None)
    yield data
    store.use_folder(None)
    store.set_resolver(None)


@pytest.fixture(autouse=True)
def no_real_llm(monkeypatch):
    """Fail loudly if a test reaches the real API without opting in."""
    from ask.agent import llm_call

    def _blocked(*a, **k):
        raise AssertionError("Test tried to call the real LLM. Use the fake_llm fixture.")

    if os.getenv("ASK_LIVE_TESTS") != "1":
        monkeypatch.setattr(llm_call, "get_client", _blocked)
    monkeypatch.setattr(llm_call, "_use_chat_api", False)
    llm_call._NEEDS_NO_REASONING.clear()


# ── sample material ────────────────────────────────────────────────────────

TRAIN_PY = '''"""Training entry point."""
from sklearn.ensemble import RandomForestClassifier

MIN_ACCURACY = 0.85


def passes_quality_gate(metrics, min_accuracy=MIN_ACCURACY):
    return metrics["accuracy"] >= min_accuracy


def train_model(X, y, seed=42):
    model = RandomForestClassifier(random_state=seed)
    model.fit(X, y)
    return model
'''

UTILS_PY = '''def load_config(path):
    """Read the YAML config."""
    with open(path) as fh:
        return fh.read()


def calculatePsi(expected, actual):
    return sum((a - e) for e, a in zip(expected, actual))
'''


@pytest.fixture
def sample_repo(tmp_path) -> Path:
    repo = tmp_path / "model_repo"
    (repo / "src").mkdir(parents=True)
    (repo / "tests").mkdir()
    (repo / "__pycache__").mkdir()
    (repo / "src" / "train.py").write_text(TRAIN_PY, encoding="utf-8")
    (repo / "src" / "utils.py").write_text(UTILS_PY, encoding="utf-8")
    (repo / "tests" / "test_train.py").write_text(
        "from src.train import passes_quality_gate\n\n"
        "def test_gate():\n    assert passes_quality_gate({'accuracy': 0.9})\n", encoding="utf-8")
    (repo / "README.md").write_text("# Model repo\n\nTrains a random forest and gates on accuracy.\n",
                                    encoding="utf-8")
    (repo / "config.yaml").write_text("min_accuracy: 0.85\nseed: 42\n", encoding="utf-8")
    (repo / "__pycache__" / "junk.py").write_text("SHOULD_BE_SKIPPED = 1\n", encoding="utf-8")
    (repo / "image.bin").write_bytes(b"\x00\x01\x02binary")
    nb = {"cells": [
        {"cell_type": "markdown", "source": ["# Exploration\n"]},
        {"cell_type": "code", "source": ["x = 1\n", "print(x)\n"],
         "outputs": [{"output_type": "stream", "text": ["1\n"]}]},
    ], "metadata": {}, "nbformat": 4, "nbformat_minor": 5}
    (repo / "explore.ipynb").write_text(json.dumps(nb), encoding="utf-8")
    return repo


@pytest.fixture
def workspace(tmp_path, sample_repo) -> Path:
    """A folder holding several projects (like the user's Repository folder)."""
    ws = tmp_path / "workspace"
    for name, readme in [("alpha", "# Alpha\n\nAlpha computes PD scores.\n"),
                         ("beta", "# Beta\n\nBeta builds lineage graphs.\n"),
                         ("gamma", "# Gamma\n\nGamma serves a RAG arena.\n")]:
        d = ws / name
        d.mkdir(parents=True)
        (d / "README.md").write_text(readme, encoding="utf-8")
        (d / "main.py").write_text(f"def run():\n    return '{name}'\n", encoding="utf-8")
        (d / "requirements.txt").write_text("pandas\n", encoding="utf-8")
    return ws


def make_portfolio(n=600, seed=0) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    x1, x2 = rng.normal(size=n), rng.normal(size=n)
    p = 1 / (1 + np.exp(-(0.9 * x1 - 0.6 * x2)))
    y = (rng.random(n) < p).astype(int)
    return pd.DataFrame({
        "x1": x1, "x2": x2,
        "segment": rng.choice(list("ABC"), n),
        "default_flag": y,
        "pd_score": np.clip(p + rng.normal(0, 0.05, n), 0, 1),
        "sample": np.where(rng.random(n) < 0.7, "train", "test"),
    })


@pytest.fixture
def portfolio_df() -> pd.DataFrame:
    return make_portfolio()


@pytest.fixture
def portfolio_csv(tmp_path, portfolio_df) -> Path:
    p = tmp_path / "portfolio.csv"
    portfolio_df.to_csv(p, index=False)
    return p


@pytest.fixture
def docx_file(tmp_path) -> Path:
    import docx
    d = docx.Document()
    d.add_heading("Model Policy", 1)
    d.add_paragraph("Models must be validated annually.")
    d.add_heading("Scope", 2)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Tier", "Frequency"
    t.cell(1, 0).text, t.cell(1, 1).text = "1", "12 months"
    p = tmp_path / "policy.docx"
    d.save(p)
    return p


@pytest.fixture
def repo_registry(sample_repo) -> SourceRegistry:
    reg = SourceRegistry()
    reg.load(sample_repo)
    return reg


# ── fake LLM (Responses API shape) ─────────────────────────────────────────

@dataclass
class FakeItem:
    type: str
    text: str = ""
    call_id: str = ""
    name: str = ""
    arguments: str = "{}"

    def model_dump(self, exclude_none=True):
        if self.type == "function_call":
            return {"type": "function_call", "call_id": self.call_id, "name": self.name,
                    "arguments": self.arguments}
        return {"type": "message", "role": "assistant",
                "content": [{"type": "output_text", "text": self.text}]}


@dataclass
class FakeResponse:
    output: list
    output_text: str = ""


def tool_call(name: str, call_id: str = "c1", **args) -> FakeItem:
    return FakeItem("function_call", call_id=call_id, name=name, arguments=json.dumps(args))


def say(text: str) -> FakeResponse:
    return FakeResponse([FakeItem("message", text=text)], output_text=text)


def calls(*items: FakeItem) -> FakeResponse:
    return FakeResponse(list(items), output_text="")


@dataclass
class FakeLLM:
    """Replays scripted responses and records every request."""
    script: list = field(default_factory=list)
    requests: list = field(default_factory=list)

    def __call__(self, **kwargs):
        self.requests.append(copy.deepcopy(kwargs))   # inputs are mutated later
        if not self.script:
            raise AssertionError("Fake LLM ran out of scripted responses")
        nxt = self.script.pop(0)
        return nxt(kwargs) if callable(nxt) else nxt


@pytest.fixture
def fake_llm(monkeypatch):
    from ask.agent import llm_call
    fake = FakeLLM()
    monkeypatch.setattr(llm_call, "_responses_create", fake)
    return fake
